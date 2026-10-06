#!/usr/bin/env python3
"""Script para gerar o ZIP inicial com os 17 SKUs do MVP."""

import zipfile
import json
import tempfile
import shutil
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import pyzipper

# =============================================================================
# Dados dos 17 SKUs do MVP
# =============================================================================

SKUS = [
    # Chips Lisa (4)
    {"categoria": "Chips Lisa", "item": "Cebola Salsa 115g", "fornecedor": "FlexPack Brasil", "sap": "100001", "versao_arte": "v1.2 - 15/03/2026", "aprovado_por": "João Silva - Eng. Alimentos"},
    {"categoria": "Chips Lisa", "item": "Clássica 70g", "fornecedor": "FlexPack Brasil", "sap": "100002", "versao_arte": "v1.0 - 10/01/2026", "aprovado_por": "Maria Santos - Eng. Alimentos"},
    {"categoria": "Chips Lisa", "item": "Queijo 80g", "fornecedor": "Embalflex", "sap": "100003", "versao_arte": "v1.1 - 20/02/2026", "aprovado_por": "Pedro Costa - Eng. Alimentos"},
    {"categoria": "Chips Lisa", "item": "Churrasco 70g", "fornecedor": "Embalflex", "sap": "100004", "versao_arte": "v1.0 - 05/01/2026", "aprovado_por": "Ana Lima - Eng. Alimentos"},
    
    # Chips Ondulada (4)
    {"categoria": "Chips Ondulada", "item": "Pimenta Mexicana 115g", "fornecedor": "Grafica São Paulo", "sap": "100005", "versao_arte": "v1.3 - 01/04/2026", "aprovado_por": "João Silva - Eng. Alimentos"},
    {"categoria": "Chips Ondulada", "item": "Picanha 115g", "fornecedor": "Grafica São Paulo", "sap": "100006", "versao_arte": "v1.2 - 15/03/2026", "aprovado_por": "Maria Santos - Eng. Alimentos"},
    {"categoria": "Chips Ondulada", "item": "Clássica 68g", "fornecedor": "FlexPack Brasil", "sap": "100007", "versao_arte": "v1.0 - 10/01/2026", "aprovado_por": "Pedro Costa - Eng. Alimentos"},
    {"categoria": "Chips Ondulada", "item": "Clássica 115g", "fornecedor": "FlexPack Brasil", "sap": "100008", "versao_arte": "v1.1 - 20/02/2026", "aprovado_por": "Ana Lima - Eng. Alimentos"},
    
    # Palha Extrafina (4)
    {"categoria": "Palha Extrafina", "item": "Tradicional 100g", "fornecedor": "PalhaPack", "sap": "100009", "versao_arte": "v1.0 - 01/02/2026", "aprovado_por": "João Silva - Eng. Alimentos"},
    {"categoria": "Palha Extrafina", "item": "Tradicional 400g", "fornecedor": "PalhaPack", "sap": "100010", "versao_arte": "v1.0 - 01/02/2026", "aprovado_por": "Maria Santos - Eng. Alimentos"},
    {"categoria": "Palha Extrafina", "item": "Sal Marinho 100g", "fornecedor": "PalhaPack", "sap": "100011", "versao_arte": "v1.1 - 15/03/2026", "aprovado_por": "Pedro Costa - Eng. Alimentos"},
    {"categoria": "Palha Extrafina", "item": "Ervas Finas 100g", "fornecedor": "PalhaPack", "sap": "100012", "versao_arte": "v1.0 - 01/02/2026", "aprovado_por": "Ana Lima - Eng. Alimentos"},
    
    # Palha Tradicional (5)
    {"categoria": "Palha Tradicional", "item": "Tradicional 100g", "fornecedor": "PalhaPack", "sap": "100013", "versao_arte": "v1.2 - 10/03/2026", "aprovado_por": "João Silva - Eng. Alimentos"},
    {"categoria": "Palha Tradicional", "item": "Tradicional 200g", "fornecedor": "PalhaPack", "sap": "100014", "versao_arte": "v1.1 - 20/02/2026", "aprovado_por": "Maria Santos - Eng. Alimentos"},
    {"categoria": "Palha Tradicional", "item": "Tradicional 400g", "fornecedor": "PalhaPack", "sap": "100015", "versao_arte": "v1.0 - 01/02/2026", "aprovado_por": "Pedro Costa - Eng. Alimentos"},
    {"categoria": "Palha Tradicional", "item": "Tradicional 800g", "fornecedor": "PalhaPack", "sap": "100016", "versao_arte": "v1.0 - 01/02/2026", "aprovado_por": "Ana Lima - Eng. Alimentos"},
    {"categoria": "Palha Tradicional", "item": "Cebola & Alho 200g", "fornecedor": "PalhaPack", "sap": "100017", "versao_arte": "v1.0 - 15/04/2026", "aprovado_por": "João Silva - Eng. Alimentos"},
]


DIMENSOES = {
    "Chips Lisa": {"largura": "180mm", "comprimento": "320mm", "espessura": "55µm", "estrutura": "BOPP/PE", "tratamento": "Corona", "cores": "4C + 1 Especial", "processo": "Flexografia", "faca": "FAC-001", "selagem": "140°C / 0.3MPa / 0.8s"},
    "Chips Ondulada": {"largura": "190mm", "comprimento": "340mm", "espessura": "60µm", "estrutura": "PET/PE", "tratamento": "Primer", "cores": "5C + 1 Especial", "processo": "Rotogravura", "faca": "FAC-002", "selagem": "150°C / 0.35MPa / 1.0s"},
    "Palha Extrafina": {"largura": "150mm", "comprimento": "280mm", "espessura": "45µm", "estrutura": "BOPP/PE", "tratamento": "Corona", "cores": "3C", "processo": "Flexografia", "faca": "FAC-003", "selagem": "135°C / 0.25MPa / 0.6s"},
    "Palha Tradicional": {"largura": "160mm", "comprimento": "300mm", "espessura": "50µm", "estrutura": "BOPP/PE", "tratamento": "Corona", "cores": "4C", "processo": "Flexografia", "faca": "FAC-004", "selagem": "138°C / 0.28MPa / 0.7s"},
}


# =============================================================================
# Funções auxiliares
# =============================================================================

def slug(cat: str, item: str) -> str:
    import unicodedata, re
    texto = f"{cat}_{item}"
    texto = unicodedata.normalize('NFKD', texto).encode('ascii', 'ignore').decode('ascii')
    texto = texto.lower()
    texto = re.sub(r'[^a-z0-9]+', '_', texto).strip('_')
    return texto


def criar_docx_template(categoria: str, item: str, dados: dict, dim: dict) -> Document:
    """Cria o DOCX da Descrição Técnica seguindo o template especificado."""
    doc = Document()
    
    # Estilo padrão
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)
    
    # Título 1 - Nome do Produto
    h1 = doc.add_heading(f'{categoria} {item}', level=1)
    h1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    # --- Identificação ---
    doc.add_heading('Identificação', level=2)
    
    identificacao = [
        ("Código SAP", dados["sap"]),
        ("Fornecedor", dados["fornecedor"]),
        ("Versão da Arte", dados["versao_arte"]),
        ("Data Aprovação", dados["versao_arte"].split(" - ")[-1] if " - " in dados["versao_arte"] else ""),
        ("Aprovado por", dados["aprovado_por"]),
    ]
    
    for label, value in identificacao:
        p = doc.add_paragraph()
        run_label = p.add_run(f"{label}: ")
        run_label.bold = True
        run_label.font.size = Pt(11)
        run_value = p.add_run(value)
        run_value.font.size = Pt(11)
    
    # --- Especificações Técnicas ---
    doc.add_heading('Especificações Técnicas', level=2)
    
    # Dimensões
    doc.add_heading('Dimensões', level=3)
    dims = [
        ("Largura", dim["largura"]),
        ("Comprimento", dim["comprimento"]),
        ("Espessura", dim["espessura"]),
    ]
    for label, value in dims:
        p = doc.add_paragraph(style='List Bullet')
        run_label = p.add_run(f"{label}: ")
        run_label.bold = True
        p.add_run(value)
    
    # Material
    doc.add_heading('Material', level=3)
    mats = [
        ("Estrutura", dim["estrutura"]),
        ("Tratamento", dim["tratamento"]),
    ]
    for label, value in mats:
        p = doc.add_paragraph(style='List Bullet')
        run_label = p.add_run(f"{label}: ")
        run_label.bold = True
        p.add_run(value)
    
    # Impressão
    doc.add_heading('Impressão', level=3)
    imps = [
        ("Cores", dim["cores"]),
        ("Processo", dim["processo"]),
    ]
    for label, value in imps:
        p = doc.add_paragraph(style='List Bullet')
        run_label = p.add_run(f"{label}: ")
        run_label.bold = True
        p.add_run(value)
    
    # Acabamento
    doc.add_heading('Acabamento', level=3)
    acabs = [
        ("Faca / Corte", dim["faca"]),
        ("Selagem", dim["selagem"]),
    ]
    for label, value in acabs:
        p = doc.add_paragraph(style='List Bullet')
        run_label = p.add_run(f"{label}: ")
        run_label.bold = True
        p.add_run(value)
    
    # --- Validação Visual ---
    doc.add_heading('Validação Visual', level=2)
    
    p = doc.add_paragraph(style='List Bullet')
    run_label = p.add_run("Referência: ")
    run_label.bold = True
    p.add_run("Conferencia Visual.pdf")
    
    p = doc.add_paragraph(style='List Bullet')
    run_label = p.add_run("Observações: ")
    run_label.bold = True
    p.add_run("Conferir cores com padrão mestre aprovado. Verificar registro de impressão e corte.")
    
    return doc


def criar_pdf_placeholder(categoria: str, item: str, output_path: Path, dim: dict = None):
    """Cria um PDF placeholder simples usando reportlab se disponível, senão cria vazio."""
    try:
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.lib.colors import HexColor
        
        c = canvas.Canvas(str(output_path), pagesize=A4)
        width, height = A4
        
        # Página 1 - Arte técnica
        c.setFillColor(HexColor("#0056A0"))
        c.rect(0, height - 60, width, 60, fill=1, stroke=0)
        c.setFillColor(HexColor("#FFFFFF"))
        c.setFont("Helvetica-Bold", 24)
        c.drawString(30*mm, height - 45*mm, f"{categoria} {item}")
        c.setFont("Helvetica", 14)
        c.drawString(30*mm, height - 55*mm, "Arte Técnica - Frente e Verso (Placeholder)")
        
        # Caixas representando frente/verso
        c.setStrokeColor(HexColor("#0056A0"))
        c.setLineWidth(2)
        c.rect(30*mm, height - 180*mm, 70*mm, 100*mm, fill=0)
        c.setFont("Helvetica", 10)
        c.drawCentredString(65*mm, height - 140*mm, "FRENTE")
        
        c.rect(130*mm, height - 180*mm, 70*mm, 100*mm, fill=0)
        c.drawCentredString(165*mm, height - 140*mm, "VERSO")
        
        # Marcas de registro
        c.setFont("Helvetica", 8)
        c.drawString(30*mm, height - 200*mm, "● Marcas de registro | ■ Guias de corte | ▲ Barras de cor")
        
        # Páginas adicionais para Chips
        if "Chips" in categoria and dim:
            # Página 2 - Especificações faca
            c.showPage()
            c.setFont("Helvetica-Bold", 16)
            c.drawString(30*mm, height - 30*mm, "Especificações de Faca e Corte")
            c.setFont("Helvetica", 11)
            y = height - 60*mm
            for k, v in dim.items():
                c.drawString(30*mm, y, f"{k}: {v}")
                y -= 12*mm
            
            # Página 3 - Guia de cores
            c.showPage()
            c.setFont("Helvetica-Bold", 16)
            c.drawString(30*mm, height - 30*mm, "Guia de Cores e Registro")
            c.setFont("Helvetica", 11)
            c.drawString(30*mm, height - 60*mm, "Cores Pantone / CMYK conforme aprovação de arte.")
            c.drawString(30*mm, height - 75*mm, "Verificar densidade de tinta e registro entre cores.")
        
        c.save()
        return True
    except ImportError:
        # Fallback: criar PDF vazio mínimo
        output_path.write_bytes(b"%PDF-1.4\n%%EOF\n")
        return False


def criar_checklist_json(item_id: str, output_path: Path):
    """Cria checklist.json inicial vazio."""
    data = {
        "item_id": item_id,
        "historico": []
    }
    output_path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")


def main():
    import argparse
    
    parser = argparse.ArgumentParser(description="Gerar ZIP inicial do Explorador de Arquivos Zip")
    parser.add_argument("-o", "--output", default="Itens.zip", help="Arquivo ZIP de saída")
    parser.add_argument("-p", "--password", default="ZipExplorer2026!", help="Senha AES-256 para o ZIP")
    args = parser.parse_args()
    
    output_zip = Path(args.output).resolve()
    senha = args.password
    
    print(f"Gerando ZIP inicial: {output_zip}")
    print(f"Senha: {senha}")
    print(f"Total de SKUs: {len(SKUS)}")
    
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        
        for sku in SKUS:
            cat = sku["categoria"]
            item = sku["item"]
            item_id = slug(cat, item)
            dim = DIMENSOES[cat]
            
            # Pasta do item
            item_dir = tmp / cat / item
            item_dir.mkdir(parents=True)
            
            print(f"  Criando: {cat}/{item}")
            
            # 1. Descrição Técnica.docx
            doc = criar_docx_template(cat, item, sku, dim)
            doc.save(item_dir / "Descricao Tecnica.docx")
            
            # 2. Conferencia Visual.pdf
            criar_pdf_placeholder(cat, item, item_dir / "Conferencia Visual.pdf", dim)
            
            # 3. checklist.json
            criar_checklist_json(item_id, item_dir / "checklist.json")
        
        # Criar ZIP com criptografia AES-256
        print("\nCompactando e criptografando (AES-256)...")
        with pyzipper.AESZipFile(
            output_zip, 'w',
            compression=pyzipper.ZIP_DEFLATED,
            encryption=pyzipper.WZ_AES
        ) as zf:
            zf.setpassword(senha.encode('utf-8'))
            zf.pwd = senha.encode('utf-8')
            
            for file_path in tmp.rglob("*"):
                if file_path.is_file():
                    arcname = file_path.relative_to(tmp).as_posix()
                    zf.write(file_path, arcname)
        
        # Verificar
        size_mb = output_zip.stat().st_size / (1024 * 1024)
        print(f"\n[OK] ZIP criado com sucesso!")
        print(f"  Arquivo: {output_zip}")
        print(f"  Tamanho: {size_mb:.2f} MB")
        print(f"  Senha: {senha}")
        
        # Testar
        print("\nTestando integridade...")
        try:
            with pyzipper.AESZipFile(output_zip, 'r') as zf:
                zf.setpassword(senha.encode('utf-8'))
                bad = zf.testzip()
                if bad:
                    print(f"  [ERRO] Arquivo corrompido: {bad}")
                else:
                    print(f"  [OK] Integridade OK - {len(zf.namelist())} arquivos")
        except Exception as e:
            print(f"  [ERRO] Erro no teste: {e}")


if __name__ == "__main__":
    main()